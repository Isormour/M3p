from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ARENAS = ROOT / "arenas"
OUTPUT = ROOT.parent / "M3P_Bioms.docx"


BIOMES = [
    {
        "name": "Dungeon",
        "status": "Obecna arena i punkt odniesienia",
        "palette": "grafitowy kamień, chłodny błękit, ciepłe pochodnie",
        "image": ARENAS / "arena-dungeon.png",
        "identity": "Monumentalne kamienne arkady, kraty, łańcuchy i wilgotne mury tworzą czytelną, zamkniętą przestrzeń.",
        "gameplay": "Środek pozostaje spokojny i kontrastowy. Ognie po bokach prowadzą wzrok, ale nie konkurują z postaciami ani planszą walki.",
        "implementation": "Zachować obecny układ jako wzorzec skali, perspektywy, wysokości platformy i gęstości detalu dla pozostałych biomów.",
    },
    {
        "name": "Lava",
        "status": "Propozycja areny",
        "palette": "czarny bazalt, czerwień żaru, intensywny pomarańcz",
        "image": ARENAS / "arena-lava.png",
        "identity": "Grota wulkaniczna z kaskadami lawy, ostrymi skałami i rozgrzanymi pęknięciami w kamieniu.",
        "gameplay": "Najjaśniejsze elementy pozostają na obrzeżach i w tle. Platforma walki jest ciemniejsza, aby sylwetki oraz efekty umiejętności były czytelne.",
        "implementation": "Animować przepływ lawy, drobne iskry i falowanie gorącego powietrza. Ograniczyć bloom przy centralnej strefie walki.",
    },
    {
        "name": "Mountain",
        "status": "Nowa koncepcja areny",
        "palette": "łupkowa szarość, śnieżny błękit, ciepłe światło świtu",
        "image": ARENAS / "arena-mountain.png",
        "identity": "Kamienny taras wysoko ponad chmurami, otoczony ruinami, mostami i ostrymi szczytami.",
        "gameplay": "Szeroka, pusta posadzka daje dużo miejsca dla bohatera i przeciwnika. Dalsze plany budują skalę bez zagęszczania pola walki.",
        "implementation": "Dodać delikatny pył śnieżny, ruch chmur i wiatr poruszający sztandarami. Jasność śniegu utrzymać poniżej efektów walki.",
    },
    {
        "name": "Forest",
        "status": "Nowa koncepcja areny",
        "palette": "głęboka zieleń, mech, ciemne drewno, akcenty turkusu",
        "image": ARENAS / "arena-forest.png",
        "identity": "Starożytna leśna polana obramowana korzeniami, omszałymi kamieniami i subtelną magią natury.",
        "gameplay": "Centralny krąg ma spokojną fakturę i umiarkowany kontrast. Jasne promienie oraz runy pozostają poza głównymi pozycjami postaci.",
        "implementation": "Wprowadzić wolne drobiny światła, ruch liści i słaby puls run. Unikać gęstych cząstek przed sylwetkami walczących.",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, name="Arial", size=10.5, bold=False, color="000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_meta_paragraph(doc, status, palette):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(status.upper())
    set_run_font(r, size=8.5, bold=True, color="555555")
    r = p.add_run("    PALETA  ")
    set_run_font(r, size=8.5, bold=True, color="777777")
    r = p.add_run(palette)
    set_run_font(r, size=8.5, color="555555")


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(label + " ")
    set_run_font(r, size=8.7, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=8.7, color="333333")
    return p


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(34)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    if title_style._element.pPr is not None:
        style_border = title_style._element.pPr.find(qn("w:pBdr"))
        if style_border is not None:
            title_style._element.pPr.remove(style_border)

    h1 = styles["Heading 1"]
    h1.font.name = "Arial"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(4)
    title.add_run("M3P Biomy Areny Walki")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("Etap 1 koncepcje środowiska")
    set_run_font(r, name="Arial", size=16, color="555555")

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(14)
    intro.paragraph_format.line_spacing = 1.12
    r = intro.add_run(
        "Dokument definiuje cztery kierunki wizualne aren walki. Dungeon stanowi obecny punkt odniesienia dla perspektywy, skali platformy i stylizowanych materiałów. Lava, Mountain i Forest rozwijają ten sam język wizualny, zachowując czytelny środek sceny dla postaci, planszy oraz efektów umiejętności."
    )
    set_run_font(r, size=11.5, color="2F2F2F")

    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [Inches(1.45), Inches(2.25), Inches(3.0), Inches(3.25)]
    headers = ["Biom", "Status", "Paleta", "Rozpoznawalny motyw"]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "2F3338")
        set_cell_margins(cell)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(headers[idx])
        set_run_font(r, size=9.5, bold=True, color="FFFFFF")

    motifs = {
        "Dungeon": "arkady, kraty i łańcuchy",
        "Lava": "kaskady lawy i bazalt",
        "Mountain": "szczyty ponad chmurami",
        "Forest": "korzenie, mech i runy",
    }
    for row_index, biome in enumerate(BIOMES, start=1):
        cells = table.add_row().cells
        values = [biome["name"], biome["status"], biome["palette"], motifs[biome["name"]]]
        for idx, (cell, value) in enumerate(zip(cells, values)):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=135, bottom=135)
            set_cell_borders(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F3F5F7")
            p = cell.paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=9.5, bold=(idx == 0), color="202020")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(14)
    note.paragraph_format.space_after = Pt(0)
    r = note.add_run("Wspolna zasada ")
    set_run_font(r, size=10.5, bold=True)
    r = note.add_run(
        "Największy kontrast i ruch są kierowane na obrzeża kadru. Centralna platforma pozostaje spokojniejsza, aby arena wspierała walkę zamiast konkurować z interfejsem i VFX."
    )
    set_run_font(r, size=10.5, color="333333")

    for biome in BIOMES:
        doc.add_page_break()
        heading = doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.space_before = Pt(0)
        heading.paragraph_format.space_after = Pt(1)
        heading.add_run(biome["name"])
        add_meta_paragraph(doc, biome["status"], biome["palette"])

        image_p = doc.add_paragraph()
        image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_p.paragraph_format.space_before = Pt(0)
        image_p.paragraph_format.space_after = Pt(3)
        image_run = image_p.add_run()
        inline = image_run.add_picture(str(biome["image"]), width=Inches(8.85))
        doc_pr = inline._inline.docPr
        doc_pr.set("descr", f"Mockup areny walki biomu {biome['name']}")

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(4)
        r = caption.add_run(f"Mockup areny walki  {biome['name']}")
        set_run_font(r, size=8, color="666666")

        add_labeled_paragraph(doc, "Kierunek", biome["identity"])
        add_labeled_paragraph(doc, "Czytelność walki", biome["gameplay"])
        add_labeled_paragraph(doc, "Implementacja", biome["implementation"])

    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run("M3P Bioms  Areny walki")
        set_run_font(r, size=8, color="777777")

    doc.core_properties.title = "M3P Biomy Areny Walki"
    doc.core_properties.subject = "Koncepcje aren walki dla biomow Lava Dungeon Mountain Forest"
    doc.core_properties.author = "M3P"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
